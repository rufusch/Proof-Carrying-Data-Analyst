import io
import json
import time
from concurrent.futures import ThreadPoolExecutor
from uuid import uuid4

import pytest
from sqlalchemy import select

from backend.common import digest, canonical
from backend.contract import validate, DOCUMENT
from backend.planner import Decision, Query, Join
from backend.sandbox import SandboxUnavailable
from backend.store import resources, jobs, events
from conftest import dataset, analysis


def test_scalar_evidence_rerun_and_restart(system):
    client, worker, store, settings, _ = system
    did = dataset(system)
    aid = analysis(system, did)
    result = client.get(f"/api/v1/analyses/{aid}/result").json()
    validate("AnalysisResult", result)
    assert result["outcome"] == "answered"
    assert result["metrics"][0]["value"] == 35
    assert result["verification"]["status"] == "passed"
    evidence = client.get(f"/api/v1/analyses/{aid}/evidence").json()
    validate("EvidencePackage", evidence)
    assert result["metrics"][0]["evidence_refs"][0] == evidence["claims"][0]["id"]
    code = evidence["code_artifacts"][0]
    assert code["sha256"] == digest(code["source"].encode())
    assert len(evidence["claims"][0]["source_locators"][0]["file_sha256"]) == 64
    from backend.api import create_app
    from fastapi.testclient import TestClient
    with TestClient(create_app(settings)) as refreshed:
        assert refreshed.get(f"/api/v1/analyses/{aid}/result").json() == result
    rerun = client.post(f"/api/v1/analyses/{aid}/reruns").json()
    assert rerun["id"] != aid and rerun["parent_analysis_id"] == aid
    worker.run_once()
    assert client.get(f"/api/v1/analyses/{aid}/evidence").json() == evidence
    assert client.get(f'/api/v1/analyses/{rerun["id"]}/result').json()["verification"]["output_hash"] == result["verification"]["output_hash"]


def test_grouped_table_chart_and_multifile(system):
    did = dataset(system, [("files", ("a.csv", b"region,amount\nWest,10\nWest,20\nEast,5\n")), ("files", ("b.csv", b"key,label\n1,other\n"))])
    aid = analysis(system, did, "sum amount by region")
    result = system[0].get(f"/api/v1/analyses/{aid}/result").json()
    assert result["outcome"] == "answered"
    assert result["tables"][0]["rows"] == [{"group_value": "East", "value": 5}, {"group_value": "West", "value": 30}]
    assert result["charts"][0]["evidence_refs"]


def test_clarification_resume_stale_and_replay(system):
    client, worker, *_ = system
    did = dataset(system, [("files", ("a.csv", b"amount\n10\n")), ("files", ("b.csv", b"amount\n20\n"))])
    aid = analysis(system, did)
    current = client.get(f"/api/v1/analyses/{aid}").json()
    assert current["status"] == "needs_clarification"
    question = current["clarification"]
    body = {"clarification_id": question["id"], "choice_id": question["choices"][1]["id"]}
    endpoint = f"/api/v1/analyses/{aid}/clarifications"
    assert client.post(endpoint, json={**body, "text": "both"}).status_code == 422
    assert client.post(endpoint, json={**body, "choice_id": "missing"}).status_code == 422
    headers = {"Idempotency-Key": "clarification-key"}
    accepted = client.post(endpoint, json=body, headers=headers)
    assert accepted.status_code == 202
    worker.run_once()
    assert client.post(endpoint, json=body, headers=headers).json() == accepted.json()
    assert client.post(endpoint, json=body).status_code == 409
    assert client.get(f"/api/v1/analyses/{aid}/result").json()["metrics"][0]["value"] == 20


def test_idempotency_all_post_routes_and_conflicts(system):
    client, worker, store, *_ = system
    files = [("files", ("a.csv", b"amount\n4\n"))]
    headers = {"Idempotency-Key": "same-request-key"}
    first = client.post("/api/v1/datasets", files=files, headers=headers)
    assert client.post("/api/v1/datasets", files=files, headers=headers).json() == first.json()
    assert client.post("/api/v1/datasets", files=[("files", ("a.csv", b"amount\n5\n"))], headers=headers).status_code == 409
    worker.run_once()
    body = {"dataset_id": first.json()["id"], "question": "sum amount"}
    one = client.post("/api/v1/analyses", json=body, headers=headers)
    two = client.post("/api/v1/analyses", json=body, headers=headers)
    assert one.json() == two.json()
    assert client.post("/api/v1/analyses", json={**body, "question": "mean amount"}, headers=headers).status_code == 409
    aid = one.json()["id"]
    cancelled = client.post(f"/api/v1/analyses/{aid}/cancel", headers=headers)
    assert client.post(f"/api/v1/analyses/{aid}/cancel", headers=headers).json() == cancelled.json()
    worker.run_once()
    assert client.get(f"/api/v1/analyses/{aid}").json()["status"] == "cancelled"
    rerun = client.post(f"/api/v1/analyses/{aid}/reruns", headers=headers)
    assert client.post(f"/api/v1/analyses/{aid}/reruns", headers=headers).json() == rerun.json()


def test_invalid_inputs_error_envelope_and_limits(system):
    client, worker, store, settings, _ = system
    for response in [client.get("/api/v1/datasets/nope"), client.get(f"/api/v1/datasets/{uuid4()}"),
        client.post("/api/v1/analyses", content="{broken", headers={"content-type": "application/json"}),
        client.post("/api/v1/analyses", json={"question": "hello"}),
        client.post("/api/v1/datasets", files=[("files", ("image.png", b"image"))]), client.get("/missing")]:
        assert response.status_code >= 400
        validate("ErrorEnvelope", response.json())
        assert response.headers["x-request-id"] == response.json()["error"]["request_id"]
    settings.max_file_bytes = 10
    assert client.post("/api/v1/datasets", files=[("files", ("a.csv", b"a\n" + b"x" * 20))]).status_code == 413
    assert not list(settings.data_dir.glob("upload-*"))


def test_not_ready_and_cancel_terminal(system):
    client, worker, *_ = system
    did = client.post("/api/v1/datasets", files=[("files", ("a.csv", b"a\n1\n"))]).json()["id"]
    assert client.get(f"/api/v1/datasets/{did}/profile").status_code == 409
    assert client.post("/api/v1/analyses", json={"dataset_id": did, "question": "sum a"}).status_code == 409
    worker.run_once()
    aid = analysis(system, did, "sum a")
    assert client.post(f"/api/v1/analyses/{aid}/cancel").status_code == 409


def test_three_failed_verifications_refuse_with_evidence(system):
    did = dataset(system)
    system[4].corrupt = True
    aid = analysis(system, did)
    result = system[0].get(f"/api/v1/analyses/{aid}/result").json()
    assert result["outcome"] == "refused" and result["metrics"] == []
    assert result["verification"]["status"] == "failed"
    evidence = system[0].get(f"/api/v1/analyses/{aid}/evidence").json()
    assert evidence["execution"]["attempt_count"] == 3
    assert len(evidence["checks"]) == 15 and evidence["claims"] == []


def test_precise_unsupported_refusal(system):
    did = dataset(system)
    for question in ["Why did revenue grow?", "sum amount and forecast next year", "run SELECT * FROM read_csv('/etc/passwd')"]:
        aid = analysis(system, did, question)
        result = system[0].get(f"/api/v1/analyses/{aid}/result").json()
        assert result["outcome"] == "refused"
        assert result["verification"]["status"] == "not_run"


def test_worker_failure_and_readiness(system):
    client, worker, *_ = system
    assert client.get("/api/v1/health").status_code == 200
    assert client.get("/api/v1/ready").status_code == 503
    did = dataset(system)
    def fail(*args):
        raise SandboxUnavailable("sensitive /server/path secret")
    worker.sandbox.run = fail
    aid = analysis(system, did)
    current = client.get(f"/api/v1/analyses/{aid}").json()
    assert current["status"] == "failed" and current["failure"]["retryable"]
    assert "secret" not in canonical(current)


def test_sse_snapshot_and_reconnect(system):
    did = dataset(system)
    aid = analysis(system, did)
    client, _, store, *_ = system
    with store.tx() as conn:
        first = conn.execute(select(events.c.seq).where(events.c.resource_id == aid).order_by(events.c.seq)).scalar()
    response = client.get(f"/api/v1/analyses/{aid}/events", headers={"Last-Event-ID": str(first)})
    assert response.headers["content-type"].startswith("text/event-stream")
    assert response.text.index("event: snapshot") < response.text.index("event: stage.changed")
    assert "event: analysis.terminal" in response.text
    for event in response.text.strip().split("\n\n"):
        assert event.startswith("id: ")
        json.loads(event.split("data: ")[1])


def test_expiry_deletes_files_events_and_jobs(system):
    did = dataset(system)
    aid = analysis(system, did)
    client, worker, store, settings, _ = system
    with store.tx() as conn:
        conn.execute(resources.update().values(expires="2000-01-01T00:00:00Z"))
    assert client.get(f"/api/v1/datasets/{did}").status_code == 404
    assert client.get(f"/api/v1/analyses/{aid}").status_code == 404
    worker.cleanup()
    assert not (settings.data_dir / did).exists()
    with store.tx() as conn:
        assert conn.execute(select(events)).first() is None
        assert conn.execute(select(jobs)).first() is None


def test_xlsx_multiple_sheets_and_warnings(system):
    from openpyxl import Workbook
    book = Workbook()
    sheet = book.active
    sheet.title = "sales"
    sheet.append(["region", "amount", "date", "currency"])
    sheet.append(["West", 10, "01/02/2025", "USD 10"])
    sheet.append(["West", None, "02/03/2025", "EUR 10"])
    second = book.create_sheet("keys")
    second.append(["id", "label"])
    second.append([1, "one"])
    output = io.BytesIO()
    book.save(output)
    did = dataset(system, [("files", ("workbook.xlsx", output.getvalue()))])
    profile = system[0].get(f"/api/v1/datasets/{did}/profile").json()
    assert len(profile["tables"]) == 2
    assert {w["code"] for w in profile["warnings"]} >= {"NULL_VALUES", "AMBIGUOUS_DATES", "CURRENCY_TEXT"}


def test_concurrent_idempotency(system):
    did = dataset(system)
    client = system[0]
    def submit(_):
        return client.post("/api/v1/analyses", json={"dataset_id": did, "question": "sum amount"}, headers={"Idempotency-Key": "concurrent-key"})
    with ThreadPoolExecutor(max_workers=4) as pool:
        responses = list(pool.map(submit, range(4)))
    assert all(r.status_code == 202 for r in responses)
    assert len({r.json()["id"] for r in responses}) == 1


def test_contract_routes_are_implemented(system):
    routes = {(r.path, method.lower()) for r in system[0].app.routes for method in (getattr(r, "methods", None) or [])}
    for path, definition in DOCUMENT["paths"].items():
        for method in definition:
            if method in {"get", "post"}:
                assert ("/api/v1" + path, method) in routes


def test_legacy_xls_all_sheets(system):
    import xlwt
    book = xlwt.Workbook()
    for name in ("first", "second"):
        sheet = book.add_sheet(name)
        sheet.write(0, 0, "amount")
        sheet.write(1, 0, 7)
    output = io.BytesIO()
    book.save(output)
    did = dataset(system, [("files", ("legacy.xls", output.getvalue()))])
    profile = system[0].get(f"/api/v1/datasets/{did}/profile").json()
    assert len(profile["tables"]) == 2
    assert any(w["code"] == "XLS_CACHED_VALUES" for w in profile["warnings"])


def test_declared_and_streamed_bundle_limits(system):
    client, _, _, settings, _ = system
    assert client.post("/api/v1/analyses", content=b"x", headers={"content-length": "99999999"}).status_code == 413
    streamed = client.post("/api/v1/analyses", content=iter([b"x" * 40000, b"x" * 40000]), headers={"content-type": "application/json"})
    assert streamed.status_code == 413
    validate("ErrorEnvelope", streamed.json())
    settings.max_dataset_bytes = 10
    response = client.post("/api/v1/datasets", files=[("files", ("a.csv", b"a\n12345\n")), ("files", ("b.csv", b"b\n12345\n"))])
    assert response.status_code == 413


def test_capacity_and_retry_after(system):
    client, _, _, settings, _ = system
    settings.max_jobs = 0
    response = client.post("/api/v1/datasets", files=[("files", ("a.csv", b"a\n1\n"))])
    assert response.status_code == 429 and response.headers["retry-after"] == "5"
    validate("ErrorEnvelope", response.json())


def test_durable_job_lease_recovery(system):
    client, worker, store, *_ = system
    did = client.post("/api/v1/datasets", files=[("files", ("a.csv", b"amount\n8\n"))]).json()["id"]
    abandoned = store.claim()
    assert abandoned is not None
    with store.tx() as conn:
        conn.execute(jobs.update().where(jobs.c.id == abandoned["id"]).values(lease=time.time() - 1))
    worker.run_once()
    assert client.get(f"/api/v1/datasets/{did}").json()["status"] == "ready"


def test_cancellation_while_execution_prevents_publication(system):
    did = dataset(system)
    client, worker, _, _, sandbox = system
    aid = client.post("/api/v1/analyses", json={"dataset_id": did, "question": "sum amount"}).json()["id"]
    original_run = sandbox.run
    def run(root, payload):
        result = original_run(root, payload)
        assert client.post(f"/api/v1/analyses/{aid}/cancel").status_code == 202
        return result
    sandbox.run = run
    worker.run_once()
    assert client.get(f"/api/v1/analyses/{aid}").json()["status"] == "cancelled"
    assert client.get(f"/api/v1/analyses/{aid}/result").json()["metrics"] == []


def test_tampered_source_fails_closed(system):
    did = dataset(system)
    client, _, _, settings, _ = system
    file = client.get(f"/api/v1/datasets/{did}").json()["files"][0]
    (settings.data_dir / did / file["id"]).write_bytes(b"tampered")
    aid = analysis(system, did)
    result = client.get(f"/api/v1/analyses/{aid}/result").json()
    assert result["outcome"] == "failed" and result["metrics"] == []


def test_clean_rerun_failure_never_answers(system):
    did = dataset(system)
    client, _, _, _, sandbox = system
    original = sandbox.run
    call_count = 0
    def inconsistent(root, payload):
        nonlocal call_count
        output = original(root, payload)
        call_count += 1
        if call_count % 2 == 0:
            output["rows"][0]["value"] += 1
            output["output_hash"] = digest(output["rows"])
        return output
    sandbox.run = inconsistent
    aid = analysis(system, did)
    result = client.get(f"/api/v1/analyses/{aid}/result").json()
    assert result["outcome"] == "refused"
    assert result["verification"]["reproducible"] is False
    assert result["metrics"] == [] and call_count == 6


def test_cleanup_recovers_orphan_uploads_without_touching_other_folders(system):
    import os
    _, worker, _, settings, _ = system
    abandoned = settings.data_dir / "upload-abandoned"
    unrelated = settings.data_dir / "user-folder"
    active = settings.data_dir / "upload-active"
    for folder in (abandoned, unrelated, active):
        folder.mkdir()
    old = time.time() - 90000
    os.utime(abandoned, (old, old))
    os.utime(unrelated, (old, old))
    worker.cleanup()
    assert not abandoned.exists() and unrelated.exists() and active.exists()
