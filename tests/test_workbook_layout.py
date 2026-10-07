import io
from openpyxl import Workbook
from conftest import dataset, analysis


def workbook(book):
    stream = io.BytesIO()
    book.save(stream)
    return [("files", ("layout.xlsx", stream.getvalue()))]


def test_empty_formatted_excel_columns_are_not_invalid_headers(system):
    book = Workbook()
    sh = book.active
    sh.append(["date", "amount"])
    sh.append(["2025-01-01", 10])
    sh.append(["2025-01-02", 20])
    sh.cell(20, 19).number_format = "0.00"
    did = dataset(system, workbook(book))
    profile = system[0].get(f"/api/v1/datasets/{did}/profile").json()
    assert profile["tables"][0]["column_count"] == 2
    assert profile["tables"][0]["row_count"] == 2
    assert any(w["code"] == "EMPTY_COLUMNS_IGNORED" for w in profile["warnings"])
    aid = analysis(system, did)
    assert system[0].get(f"/api/v1/analyses/{aid}/result").json()["metrics"][0]["value"] == 30


def test_excluded_sheets_report_reasons_and_propagate_to_answers(system):
    book = Workbook()
    book.active.title = "Transactions"
    book.active.append(["amount"])
    book.active.append([10])
    formulas = book.create_sheet("Calculated")
    formulas.append(["amount"])
    formulas.append(["=1+1"])
    summary = book.create_sheet("Summary")
    summary.append([None, "Total", 10])
    did = dataset(system, workbook(book))
    profile = system[0].get(f"/api/v1/datasets/{did}/profile").json()
    assert len(profile["tables"]) == 1
    excluded = [w for w in profile["warnings"] if w["code"] == "SHEET_EXCLUDED"]
    assert len(excluded) == 2 and any("paste calculated values" in w["message"] for w in excluded)
    aid = analysis(system, did)
    result = system[0].get(f"/api/v1/analyses/{aid}/result").json()
    assert result["metrics"][0]["value"] == 10
    assert sum(w["code"] == "SHEET_EXCLUDED" for w in result["warnings"]) == 2


def test_no_valid_sheets_returns_actionable_failure(system):
    book = Workbook()
    book.active.append(["amount"])
    book.active.append(["=1+1"])
    client, worker, *_ = system
    created = client.post("/api/v1/datasets", files=workbook(book)).json()
    worker.run_once()
    data = client.get('/api/v1/datasets/' + created['id']).json()
    assert data['status'] == 'failed'
    assert 'No usable tables' in data['failure']['message']
    assert 'Row 2 contains a formula' in data['failure']['message']


def test_populated_unnamed_columns_are_not_silently_removed(system):
    book = Workbook()
    book.active.append(["amount", None])
    book.active.append([10, "unlabelled value"])
    client, worker, *_ = system
    created = client.post("/api/v1/datasets", files=workbook(book)).json()
    worker.run_once()
    data = client.get('/api/v1/datasets/' + created['id']).json()
    assert data['status'] == 'failed' and 'without headers (2)' in data['failure']['message']


def test_currency_question_explains_column_type_and_fix(system):
    book = Workbook()
    book.active.append(["Revenue", "Quarter"])
    book.active.append(["$1,000 USD", "Q1"])
    did = dataset(system, workbook(book))
    for question in ['How much revenue did we make?', 'sum Revenue']:
        aid = analysis(system, did, question)
        result = system[0].get(f'/api/v1/analyses/{aid}/result').json()
        assert result['outcome'] == 'refused'
        assert 'Revenue' in result['narrative'] and 'numeric' in result['narrative']
        assert 'copy' in result['narrative']


def test_missing_metric_lists_available_columns(system):
    did = dataset(system)
    aid = analysis(system, did, 'What is our total profit?')
    result = system[0].get(f'/api/v1/analyses/{aid}/result').json()
    assert result['outcome'] == 'refused'
    assert 'profit' in result['narrative'] and 'amount' in result['narrative']
