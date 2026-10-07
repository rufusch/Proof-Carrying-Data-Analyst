import csv
import json
from scripts.make_test_data import generate
from backend.common import digest


def test_synthetic_duplicate_data_and_known_totals(tmp_path):
    root = generate(tmp_path)
    assert digest((root / "sales.csv").read_bytes()) == digest((root / "sales-duplicate.csv").read_bytes())
    with (root / "sales.csv").open() as handle:
        rows = list(csv.DictReader(handle))
    expected = json.loads((root / "expected.json").read_text())
    assert sum(int(r["amount"]) for r in rows) == expected["total_amount"]
    assert sum(int(r["units"]) for r in rows) == expected["total_units"]
    from openpyxl import load_workbook
    book = load_workbook(root / "sales-with-duplicate-sheet.xlsx", read_only=True)
    try:
        assert len(book.sheetnames) == 2
        assert list(book.worksheets[0].values) == list(book.worksheets[1].values)
    finally:
        book.close()
