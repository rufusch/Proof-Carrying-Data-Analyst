"""Real Docker-backed acceptance test; expected answers come independently from original cells."""
import argparse
from datetime import datetime
from decimal import Decimal
import json
from pathlib import Path
import sys
import time
from openpyxl import load_workbook
import httpx

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from backend.contract import validate


def main():
    parser=argparse.ArgumentParser(); parser.add_argument('dataset',type=Path); args=parser.parse_args()
    book=load_workbook(args.dataset,read_only=True,data_only=False)
    try:
        rows=list(book.active.iter_rows(values_only=True)); headers=rows[0]
        dcol,vcol=headers.index('Quarter'),headers.index('Revenue')
        original=[(r[dcol],Decimal(str(r[vcol]).replace('$','').replace(',',''))) for r in rows[1:] if isinstance(r[dcol],datetime) and r[vcol] is not None]
        latest=max(d for d,v in original); quarter=(latest.month-1)//3
        quarter_total=sum(v for d,v in original if d.year==latest.year and (d.month-1)//3==quarter)
        total=sum(v for d,v in original)
        previous_index=latest.year*4+quarter-1
        previous_total=sum(v for d,v in original if d.year*4+(d.month-1)//3==previous_index)
    finally:book.close()
    report={'expected_latest_quarter':f'Q{quarter+1} {latest.year}','expected_quarter_revenue':str(quarter_total),'expected_total_revenue':str(total),'analyses':[]}
    with httpx.Client(base_url='http://127.0.0.1:8010/api/v1',timeout=30) as client:
        def checked(response,schema,status=200):
            assert response.status_code==status,response.text
            return validate(schema,response.json())
        def wait(path,schema,states):
            deadline=time.monotonic()+240
            while time.monotonic()<deadline:
                data=checked(client.get(path),schema)
                if data['status'] in states:return data
                time.sleep(.5)
            raise RuntimeError('Timed out')
        deadline=time.monotonic()+90
        while client.get('/ready').status_code!=200:
            assert time.monotonic()<deadline,'Backend not ready';time.sleep(1)
        created=checked(client.post('/datasets',files=[('files',(args.dataset.name,args.dataset.read_bytes()))]),'Dataset',202); did=created['id']
        assert wait('/datasets/'+did,'Dataset',{'ready','failed'})['status']=='ready'
        for question,expected in [('what is the revenue made in the last quarter',quarter_total),('what is the total revenue',total),('Revenue?',total),('What is the change in revenue from the previous quarter to the latest quarter?',quarter_total-previous_total)]:
            created=checked(client.post('/analyses',json={'dataset_id':did,'question':question}),'Analysis',202); aid=created['id']
            state=wait('/analyses/'+aid,'Analysis',{'completed','refused','failed','needs_clarification'}); assert state['status']=='completed',state
            result=checked(client.get('/analyses/'+aid+'/result'),'AnalysisResult')
            evidence=checked(client.get('/analyses/'+aid+'/evidence'),'EvidencePackage')
            metric=next(m for m in result['metrics'] if m['id']=='metric_value')
            assert Decimal(str(metric['value']))==expected
            assert result['verification']['status']=='passed' and result['verification']['reproducible']
            assert result['metrics'][0]['unit']=='$'
            assert result['metrics'][0]['formatted_value'] in result['narrative']
            count=sum(d.year==latest.year and (d.month-1)//3==quarter for d,v in original) if 'quarter' in question else len(original)
            if 'change' in question:
                assert '−' in result['narrative'] and metric['formatted_value'] in result['narrative']
            else:
                assert f'Added {count} Revenue value' in result['narrative']
            assert 'source hashes' not in result['narrative'] and 'SQL' not in result['narrative']
            report['analyses'].append({'question':question,'result':result,'evidence':evidence})
            print(json.dumps({'question':question,'value':metric['formatted_value'],'confidence':result['confidence']['score'],'verification':result['verification']['status']}),flush=True)
        assert report['analyses'][2]['result']['confidence']['score']<report['analyses'][1]['result']['confidence']['score']
    report['status']='passed'
    path=ROOT/'.test-env/tesla-report.json';path.write_text(json.dumps(report,indent=2)+'\n',encoding='utf-8')
    print('PASS: independently derived Tesla totals, dataset-relative quarter, currency display, and reduced vague-question confidence.')


if __name__=='__main__':main()
