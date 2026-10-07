"""Adapt the local intent model using an explicitly supplied workbook schema, not its answers."""
import argparse
import hashlib
import json
from pathlib import Path
import sys
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("dataset", type=Path)
    args = parser.parse_args()
    book = load_workbook(args.dataset, read_only=True, data_only=False)
    try:
        columns = [c.value for c in next(book.active.iter_rows()) if isinstance(c.value, str) and c.value.strip()]
        if "Revenue" not in columns or "Quarter" not in columns:
            raise SystemExit("This adaptation requires Revenue and Quarter fields.")
        examples = []
        for column in [c for c in columns if c in {"Revenue", "Gross Profit", "Net Income", "Share Price", "Total Assets", "Total Liabilities"}]:
            for intent, phrase in [("sum", "total"), ("mean", "average"), ("min", "minimum"), ("max", "maximum"), ("count", "count")]:
                examples += [{"intent": intent, "text": f"{phrase} {column}"}, {"intent": intent, "text": f"what is the {phrase} {column}"}]
        examples += [{"intent": "sum", "text": q} for q in ["what is the revenue made", "what is the revenue", "show revenue", "what is the total revenue", "how much revenue did Tesla make"]]
        adapted = {"source_sha256": hashlib.sha256(args.dataset.read_bytes()).hexdigest(), "columns": columns, "examples": examples,
            "scope": "Authorized schema-based intent adaptation. Cell values and calculated answers are not memorized or included in training."}
        (ROOT / "docs/dataset-planner-training.json").write_text(json.dumps(adapted, indent=2) + "\n", encoding="utf-8")
    finally:
        book.close()
    from scripts.train_intent_model import main as train
    train()


if __name__ == "__main__":
    main()
