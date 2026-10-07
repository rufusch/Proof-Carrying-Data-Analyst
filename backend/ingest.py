"""Bounded tabular parsing, used inside the sandbox, never in the API process."""
import csv
import io
import math
import re
import zipfile
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import pandas as pd

from .common import digest, now, warning


class IngestError(ValueError):
    """Only fixed, client-safe validation messages cross the sandbox boundary."""


CURRENCY_MARKERS = {"$": "DOLLAR", "€": "EURO", "£": "POUND", "₹": "RUPEE", "USD": "USD", "EUR": "EUR", "GBP": "GBP", "INR": "INR"}


def currency_numbers(values):
    """Lossless removal of one consistent currency marker; never FX or scale conversion."""
    amount = r"[-+]?(?:0|[1-9]\d*|[1-9]\d{0,2}(?:,\d{3})+)(?:\.\d+)?"
    parsed, markers = [], set()
    for value in values:
        if value is None:
            parsed.append(None)
            continue
        if not isinstance(value, str):
            return None
        match = re.fullmatch(r"\s*(\()?\s*(USD|EUR|GBP|INR|[$€£₹])\s*(" + amount + r")\s*(\))?\s*", value)
        if not match or bool(match[1]) != bool(match[4]) or (match[1] and match[3][0] in "+-"):
            return None
        markers.add(match[2])
        number = Decimal(match[3].replace(",", "")) * (-1 if match[1] else 1)
        if abs(number) > 2**53 or not math.isfinite(float(number)) or Decimal(str(float(number))) != number:
            return None
        parsed.append(int(number) if number == number.to_integral_value() else float(number))
    return (parsed, next(iter(markers))) if len(markers) == 1 else None


def scalar(value):
    if value is None or pd.isna(value):
        return None
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if hasattr(value, "item"):
        value = value.item()
    if isinstance(value, float) and not math.isfinite(value):
        raise IngestError("Non-finite numeric values are not supported.")
    if isinstance(value, int) and abs(value) > 2**53:
        return str(value)  # preserve identifiers that JSON/JS cannot represent exactly
    if isinstance(value, str) and len(value) > 10000:
        raise IngestError("Cell text exceeds the 10,000-character limit.")
    return value


def parse_file(path, extension, max_rows, max_columns, diagnostics=None):
    if extension == ".csv":
        content = path.read_bytes().decode("utf-8-sig")
        reader = csv.reader(io.StringIO(content))
        header = next(reader, [])
        check_header(header, max_columns)
        count = 0
        for row in reader:
            if not row:
                continue
            if len(row) != len(header):
                raise IngestError("CSV rows must have the same width as the header.")
            count += 1
            if count > max_rows:
                raise IngestError("Table row limit exceeded.")
        frame = pd.read_csv(io.StringIO(content), dtype=str, keep_default_na=False, na_values=[""], nrows=max_rows + 1)
        # Infer without a float intermediary: preserve leading-zero IDs and large integers.
        for name in frame.columns:
            present = [v for v in frame[name] if not pd.isna(v)]
            converter = None
            if present and all(re.fullmatch(r"-?(?:0|[1-9]\d*)", v) and abs(int(v)) <= 2**53 for v in present):
                converter = int
            elif present and all(re.fullmatch(r"-?(?:0|[1-9]\d*)(?:\.\d+)?(?:[eE][-+]?\d+)?", v)
                    and math.isfinite(float(v)) and abs(Decimal(v)) <= 2**53
                    and Decimal(v) == Decimal(str(float(v))) for v in present):
                converter = float
            elif present and all(v.lower() in {"true", "false"} for v in present):
                converter = lambda v: v.lower() == "true"
            if converter:
                frame[name] = pd.Series([None if pd.isna(v) else converter(v) for v in frame[name]], dtype=object)
        yield None, frame
        return
    if extension == ".xlsx":
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > 10000 or sum(x.file_size for x in members) > 256 * 1024**2:
                raise IngestError("Expanded workbook exceeds the safe parsing limit.")
            if any(x.file_size > 1024**2 and x.file_size / max(x.compress_size, 1) > 1000 for x in members):
                raise IngestError("Workbook compression ratio exceeds the safe parsing limit.")
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(path.read_bytes()), read_only=True, data_only=False, keep_links=False)
        try:
            for sheet in book:
                try:
                    frame = parse_xlsx_sheet(sheet, max_rows, max_columns)
                except IngestError as exc:
                    if diagnostics is None:
                        raise
                    diagnostics.append(warning("SHEET_EXCLUDED", f'Sheet "{sheet.title}" was excluded: {exc}'))
                    continue
                yield sheet.title, frame
        finally:
            book.close()
    else:
        import xlrd
        book = xlrd.open_workbook(path, on_demand=True)
        try:
            for sheet in book.sheets():
                if sheet.nrows > max_rows + 1:
                    raise IngestError("Table row limit exceeded.")
                header = sheet.row_values(0) if sheet.nrows else []
                check_header(header, max_columns)
                rows = []
                for i in range(1, sheet.nrows):
                    row = []
                    for cell in sheet.row(i):
                        if cell.ctype == xlrd.XL_CELL_DATE:
                            value = xlrd.xldate_as_datetime(cell.value, book.datemode)
                        elif cell.ctype == xlrd.XL_CELL_ERROR:
                            raise IngestError("Workbook contains an Excel error cell.")
                        else:
                            value = None if cell.ctype in (0, 6) or cell.value == "" else cell.value
                        row.append(value)
                    rows.append(row)
                yield sheet.name, pd.DataFrame(rows, columns=header, dtype=object)
        finally:
            book.release_resources()


def parse_xlsx_sheet(sheet, max_rows, max_columns):
    if sheet.max_column and sheet.max_column > max_columns:
        raise IngestError(f"Table exceeds {max_columns} columns. Remove unused formatted columns or export the data range as CSV.")
    iterator = sheet.iter_rows()
    first = next(iterator, ())
    header = [cell.value for cell in first]
    if any(cell.data_type == "f" for cell in first):
        raise IngestError("The header contains a formula. Put plain column names in the first row.")
    present = [i for i, value in enumerate(header) if value is not None and (not isinstance(value, str) or value.strip())]
    named = [header[i] for i in present]
    check_header(named, max_columns)
    rows, empty_columns = [], set(range(len(header))) - set(present)
    for row_number, row in enumerate(iterator, 2):
        if len(row) > max_columns:
            raise IngestError(f"Row {row_number} exceeds {max_columns} columns. Export only the intended table range.")
        if any(cell.data_type == "f" for cell in row):
            raise IngestError(f"Row {row_number} contains a formula. Formulas are not evaluated; paste calculated values into a copy before uploading.")
        values = [cell.value for cell in row]
        if not any(value is not None and value != "" for value in values):
            continue
        unnamed = [i + 1 for i, value in enumerate(values) if i not in present and value is not None and value != ""]
        if unnamed:
            raise IngestError(f"Row {row_number} has data in columns without headers ({', '.join(map(str, unnamed[:8]))}). Add names in row 1, or export the intended table range as CSV.")
        empty_columns.update(i for i in range(len(values)) if i not in present)
        rows.append([values[i] if i < len(values) else None for i in present])
        if len(rows) > max_rows:
            raise IngestError(f"Table exceeds {max_rows} data rows. Split the table into smaller files.")
    frame = pd.DataFrame(rows, columns=named, dtype=object)
    frame.attrs["parsing_warnings"] = [warning("EMPTY_COLUMNS_IGNORED", f"Ignored {len(empty_columns)} completely empty formatted columns; no populated cells were removed.")] if empty_columns else []
    return frame


def check_header(header, maximum):
    if not header:
        raise IngestError("No column headers were found. Put column names in the first row, remove blank title rows, and save a copy as CSV or XLSX.")
    if len(header) > maximum:
        raise IngestError(f"The header contains {len(header)} columns; the limit is {maximum}. Remove unused formatted columns or split the table.")
    invalid = [i + 1 for i, value in enumerate(header) if not isinstance(value, str) or not value.strip() or len(value) > 200]
    if invalid:
        raise IngestError(f"Invalid headers in columns {', '.join(map(str, invalid[:8]))}. Each column needs a non-empty text name of at most 200 characters in row 1. Summary-only sheets and merged title rows are not tables.")
    if len({x.casefold() for x in header}) != len(header):
        raise IngestError("Column names must be unique, including case.")


def ingest(root, dataset, limits):
    tables, stored, warnings = [], {}, []
    total_rows = 0
    for file in dataset["files"]:
        extension = Path(file["filename"]).suffix.lower()
        path = root / file["id"]
        if digest(path.read_bytes()) != file["sha256"]:
            raise IngestError("Source file hash mismatch.")
        diagnostics = []
        for sheet, frame in parse_file(path, extension, limits["max_rows"], limits["max_columns"], diagnostics):
            total_rows += len(frame)
            if total_rows > limits["max_rows"] or len(tables) >= limits["max_tables"]:
                raise IngestError("Dataset expansion exceeds effective row/table limits.")
            id = "t_" + file["id"].replace("-", "") + "_" + str(len(tables))
            warnings.extend({**w, "table_id": id} for w in frame.attrs.get("parsing_warnings", []))
            rows = [{str(k): scalar(v) for k, v in row.items()} for row in frame.to_dict(orient="records")]
            cols = []
            for name in frame.columns:
                normalized_currency = currency_numbers([row[name] for row in rows])
                currency_warning = []
                if normalized_currency:
                    numbers, marker = normalized_currency
                    for row, number in zip(rows, numbers):
                        row[name] = number
                    currency_warning = [warning("CURRENCY_FORMAT_" + CURRENCY_MARKERS[marker], f'Removed consistent "{marker}" markers and valid thousands separators for numeric calculation. No currency conversion or million/billion scaling was applied; amounts use the file\'s original scale.', table_id=id, column=name)]
                values = [r[name] for r in rows if r[name] is not None]
                types = {type(v) for v in values}
                kind = "unknown" if not values else "boolean" if types <= {bool} else "integer" if types <= {int} else "number" if types <= {int, float} else "string"
                cw = currency_warning
                if kind == "string":
                    strings = [str(v) for v in values]
                    if any(re.fullmatch(r"[-+]?\d+(?:\.\d+)?(?:[eE][-+]?\d+)?", x) for x in strings):
                        cw.append(warning("NUMERIC_TEXT", "Numeric-looking text was preserved to avoid losing identifier formatting or numeric precision.", table_id=id, column=name))
                    if strings and all(re.fullmatch(r"\d{4}-\d{2}-\d{2}(T.*)?", x) for x in strings):
                        kind = "datetime" if any("T" in x for x in strings) else "date"
                    if any(re.search(r"\b\d{1,2}[/.-]\d{1,2}[/.-]\d{2,4}\b", x) for x in strings):
                        cw.append(warning("AMBIGUOUS_DATES", "Date order is ambiguous; values were preserved as text.", table_id=id, column=name))
                    currencies = {s for x in strings for s in re.findall(r"[$€£₹]|\b(?:USD|EUR|GBP|INR)\b", x)}
                    if currencies:
                        cw.append(warning("CURRENCY_TEXT", "Currency values remain text; confirm currency and conversion before calculation.", table_id=id, column=name))
                    if any(re.fullmatch(r"\s*[-+]?\d+(?:\.\d+)?\s*[A-Za-z%]+\s*", x) for x in strings):
                        cw.append(warning("UNIT_TEXT", "Values contain units; no unit conversion has been applied.", table_id=id, column=name))
                nulls = len(rows) - len(values)
                distinct = len({digest(x) for x in values})
                if nulls:
                    cw.append(warning("NULL_VALUES", f"{nulls} values are missing; aggregate functions exclude missing values.", table_id=id, column=name))
                cols.append({"name": name, "inferred_type": kind, "nullable": bool(nulls), "null_count": nulls,
                    "distinct_count": distinct, "sample_values": list(dict.fromkeys(values))[:10],
                    "candidate_key": bool(values) and not nulls and distinct == len(rows), "warnings": cw,
                    "min": min(values) if values and kind in {"number", "integer", "date", "datetime"} else None,
                    "max": max(values) if values and kind in {"number", "integer", "date", "datetime"} else None})
                warnings.extend(cw)
            duplicates = len(rows) - len({digest(r) for r in rows})
            if duplicates:
                warnings.append(warning("DUPLICATE_ROWS", f"{duplicates} duplicate rows are retained.", table_id=id))
            if extension == ".xls":
                warnings.append(warning("XLS_CACHED_VALUES", "Legacy Excel cells use stored values; formula freshness cannot be established.", table_id=id))
            table = {"id": id, "file_id": file["id"], "name": Path(file["filename"]).stem + (f" / {sheet}" if sheet else ""),
                "sheet_name": sheet, "row_count": len(rows), "column_count": len(cols), "duplicate_row_count": duplicates,
                "columns": cols, "sample_rows": rows[:10]}
            tables.append(table)
            stored[id] = rows
        warnings.extend({**w, "message": f'{file["filename"]}: {w["message"]}'} for w in diagnostics)
    if not tables:
        reasons = " ".join(w["message"] for w in warnings[:4])
        raise IngestError("No usable tables were found. " + reasons)
    joins = []
    for i, left in enumerate(tables):
        for right in tables[i + 1:]:
            for lc in left["columns"]:
                for rc in right["columns"]:
                    if lc["name"].casefold() == rc["name"].casefold() and lc["inferred_type"] == rc["inferred_type"] and (lc["candidate_key"] or rc["candidate_key"]):
                        joins.append({"left_table": left["id"], "left_column": lc["name"], "right_table": right["id"], "right_column": rc["name"], "confidence": 0.6, "reason": "Matching name/type and a candidate key; relationship semantics require confirmation."})
    from .dataset_vocabulary import build_vocabulary
    profile = {"dataset_id": dataset["id"], "generated_at": now(), "tables": tables, "join_candidates": joins[:100], "warnings": warnings}
    return {"profile": profile, "tables": stored, "vocabulary": build_vocabulary(profile, stored)}

